from __future__ import annotations

from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Cm, Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = Path('/mnt/data/금융권_N2SF_보안통제_실증논문_초안_v1.docx')
FIG = ROOT / 'figures'

TITLE = '금융권 N2SF 시스템·업무별 모델 기반 보안통제 설계 및 역할 기반 마이크로세그멘테이션 검증'
ENG_TITLE = 'Security Control Design and Role-Based Micro-Segmentation Validation for a System- and Business-Function-Based N2SF Model in the Financial Sector'


def set_cell_shading(cell, fill: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tcMar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    tblHeader.set(qn('w:val'), 'true')
    trPr.append(tblHeader)


def set_row_cant_split(row):
    """Prevent a table row from being split across pages."""
    trPr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement('w:cantSplit')
    trPr.append(cant_split)


def set_repeat_header_text(cell, text, bold=True):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(8.5)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fldChar1 = OxmlElement('w:fldChar'); fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText'); instrText.set(qn('xml:space'), 'preserve'); instrText.text = ' PAGE '
    fldChar2 = OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1); run._r.append(instrText); run._r.append(fldChar2)


def add_p(doc, text='', style=None, bold_prefix=None, align=None, space_after=4, first_indent=True):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.35
    if first_indent and style is None:
        p.paragraph_format.first_line_indent = Pt(9)
    if bold_prefix and text.startswith(bold_prefix):
        p.add_run(bold_prefix).bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p


def add_bullet(doc, text, level=0):
    style = 'List Bullet' if level == 0 else 'List Bullet 2'
    p = doc.add_paragraph(text, style=style)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.2
    return p


def add_number(doc, text, level=0):
    style = 'List Number' if level == 0 else 'List Number 2'
    p = doc.add_paragraph(text, style=style)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.2
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f'Heading {level}')
    p.paragraph_format.keep_with_next = True
    return p


def add_caption(doc, text):
    p = doc.add_paragraph(style='Caption')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(text)
    p.paragraph_format.space_after = Pt(5)
    return p


def add_figure(doc, filename, caption, width_cm=15.8):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(FIG / filename), width=Cm(width_cm))
    add_caption(doc, caption)


def add_table(doc, headers, rows, widths=None, font_size=8.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    set_row_cant_split(hdr)
    for i, h in enumerate(headers):
        set_repeat_header_text(hdr.cells[i], h)
        set_cell_shading(hdr.cells[i], 'D9E2F3')
        set_cell_margins(hdr.cells[i])
        if widths:
            hdr.cells[i].width = Cm(widths[i])
    for row in rows:
        new_row = table.add_row()
        set_row_cant_split(new_row)
        cells = new_row.cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                for r in p.runs:
                    r.font.size = Pt(font_size)
            if widths:
                cells[i].width = Cm(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_equation(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.font.name = 'Cambria Math'
    r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Cambria Math')
    r.font.size = Pt(10.5)
    return p


def add_placeholder(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run(f'[실험 후 입력: {text}]')
    r.bold = True
    r.font.color.rgb = RGBColor(176, 0, 32)
    return p


def configure_styles(doc: Document):
    styles = doc.styles
    normal = styles['Normal']
    normal.font.name = 'Noto Sans CJK KR'
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Noto Sans CJK KR')
    normal.font.size = Pt(9.5)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(4)

    for name, size, color in [('Heading 1', 15, '1F4E79'), ('Heading 2', 12, '1F4E79'), ('Heading 3', 10.5, '404040')]:
        st = styles[name]
        st.font.name = 'Noto Sans CJK KR'
        st._element.rPr.rFonts.set(qn('w:eastAsia'), 'Noto Sans CJK KR')
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(10 if name == 'Heading 1' else 7)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.keep_with_next = True

    cap = styles['Caption']
    cap.font.name = 'Noto Sans CJK KR'
    cap._element.rPr.rFonts.set(qn('w:eastAsia'), 'Noto Sans CJK KR')
    cap.font.size = Pt(8.5)
    cap.font.italic = False
    cap.font.color.rgb = RGBColor(64,64,64)

    for list_name in ['List Bullet','List Bullet 2','List Number','List Number 2']:
        st = styles[list_name]
        st.font.name = 'Noto Sans CJK KR'
        st._element.rPr.rFonts.set(qn('w:eastAsia'), 'Noto Sans CJK KR')
        st.font.size = Pt(9.2)

    if 'Abstract' not in styles:
        s = styles.add_style('Abstract', WD_STYLE_TYPE.PARAGRAPH)
        s.font.name = 'Noto Sans CJK KR'
        s._element.rPr.rFonts.set(qn('w:eastAsia'), 'Noto Sans CJK KR')
        s.font.size = Pt(9)
        s.paragraph_format.line_spacing = 1.25
        s.paragraph_format.space_after = Pt(3)


def main():
    doc = Document()
    configure_styles(doc)
    sec = doc.sections[0]
    sec.page_width = Cm(21.0); sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.0); sec.bottom_margin = Cm(1.8); sec.left_margin = Cm(2.1); sec.right_margin = Cm(2.1)
    sec.header_distance = Cm(0.8); sec.footer_distance = Cm(0.8)
    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = header.add_run('금융권 N2SF Track 2 연구초안')
    hr.font.size = Pt(8); hr.font.color.rgb = RGBColor(110,110,110)
    add_page_number(sec.footer.paragraphs[0])

    # Cover
    p = doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(90)
    r=p.add_run(TITLE); r.bold=True; r.font.size=Pt(22); r.font.name='Noto Sans CJK KR'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Sans CJK KR')
    p2=doc.add_paragraph(); p2.alignment=WD_ALIGN_PARAGRAPH.CENTER; p2.paragraph_format.space_before=Pt(16)
    r=p2.add_run(ENG_TITLE); r.italic=True; r.font.size=Pt(12)
    p3=doc.add_paragraph(); p3.alignment=WD_ALIGN_PARAGRAPH.CENTER; p3.paragraph_format.space_before=Pt(70)
    rr=p3.add_run('결과 입력 전 연구초안'); rr.bold=True; rr.font.size=Pt(13); rr.font.color.rgb=RGBColor(176,0,32)
    for line in ['저자: [입력]', '소속: [입력]', '작성일: 2026년 8월']:
        pp=doc.add_paragraph(); pp.alignment=WD_ALIGN_PARAGRAPH.CENTER; pp.add_run(line).font.size=Pt(10.5)
    add_p(doc, '본 문서는 선행 논문(Track 1)이 금융권 N2SF 시스템·업무별 모델을 제안하였다는 전제에서, 해당 모델의 보안통제 설계와 실험 검증을 다루는 후속 논문 초안이다. 실제 실험을 수행하지 않았으며 결과표와 해석 부분은 공란 또는 입력 표시로 남겨 두었다.', align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False, space_after=0)
    doc.add_page_break()

    # Abstract
    add_heading(doc, '요약', 1)
    abstract = (
        '금융분야의 망분리 정책은 외부 위협과 내부 업무환경을 네트워크 경계로 구분하여 보호하는 데 기여해 왔으나, '
        '클라우드, SaaS 및 생성형 인공지능의 활용 확대에 따라 업무시스템 사이의 동서 통신과 S/O 영역 간 정보이동을 세밀하게 통제할 필요가 커지고 있다. '
        '본 연구는 선행 연구가 제안한 금융권 N2SF 시스템·업무별 모델을 전제로, 가상 금융기관 A의 고객정보, 여신심사, 신용평가, 자금세탁방지 및 승인 업무에 적용할 보안통제 항목을 N2SF 정보서비스 모델 해설서의 절차에 따라 도출한다. '
        '또한 NSDI 2025에서 제안된 ZTS의 역할 기반 마이크로세그멘테이션 개념과 정책 위반률 평가 절차를 금융권 실험환경에 맞게 적용한다. '
        '비교군은 모든 S영역 서비스와 데이터베이스가 공통 내부 네트워크를 공유하고 인증 후 광범위한 접근을 허용하는 구조로 구성하며, 제안군은 업무별 컨테이너 네트워크, 정책집행지점(PEP), Open Policy Agent 기반 정책결정지점(PDP), Transfer CDS 모의 구성요소를 적용한다. '
        '평가는 정상 업무 흐름 성공률, 비인가 흐름 차단률, 업무 간 데이터베이스 직접 도달률, 침해 노드의 공격 확산 범위, S/O 전송통제 효과, 정책 결정 지연 및 감사로그 완전성을 기준으로 수행한다. '
        '본 논문 초안은 실제 측정값을 제시하지 않으며, 재현 가능한 실험환경과 결과 입력 구조를 제공함으로써 후속 실험을 위한 연구 설계를 제시한다.'
    )
    add_p(doc, abstract, style='Abstract', first_indent=True)
    add_p(doc, '주제어: 국가 망 보안체계, N2SF, 금융권 망분리, 마이크로세그멘테이션, 정책 기반 접근통제, CDS', bold_prefix='주제어:', first_indent=False)
    add_heading(doc, 'Abstract', 1)
    eng_abs = (
        'Network separation has reduced the attack surface of financial institutions by isolating internal information-processing environments from external networks. '
        'However, the growing use of cloud services, SaaS, and generative AI requires fine-grained control over east-west communications and information transfers between Sensitive and Open zones. '
        'Assuming that a preceding Track 1 study has already proposed a system- and business-function-based N2SF model for the financial sector, this study derives security requirements and controls for a hypothetical Financial Institution A by following the structure of the N2SF information-service model appendices. '
        'The study adapts the role-based micro-segmentation and policy-violation evaluation process of ZTS, presented at NSDI 2025. '
        'A flat internal-network baseline is compared with a proposed environment consisting of business-specific container networks, a policy enforcement point, an Open Policy Agent policy decision point, and a Transfer-CDS prototype. '
        'The evaluation plan measures legitimate workflow availability, unauthorized-flow blocking, direct database reachability, blast radius, S/O transfer control, policy-decision latency, and audit completeness. '
        'No experimental values are fabricated in this draft; instead, it provides a reproducible environment and result templates for subsequent execution.'
    )
    add_p(doc, eng_abs, style='Abstract', first_indent=False)
    add_p(doc, 'Keywords: N2SF, financial network separation, micro-segmentation, policy as code, cross-domain solution', bold_prefix='Keywords:', first_indent=False)

    add_heading(doc, '논문 구성', 1)
    contents = [
        '1. 서론', '2. 연구 배경 및 검증 방법론 선정', '3. 가상 금융기관 A의 보안위협 및 보안통제 도출',
        '4. 역할 기반 마이크로세그멘테이션 실험환경 구현', '5. 실험 설계 및 평가방법', '6. 실험 결과 작성 구조',
        '7. 논의', '8. 결론', '참고문헌', '부록 A. 실험 실행 절차'
    ]
    for c in contents: add_bullet(doc, c)
    doc.add_page_break()

    # 1 Introduction
    add_heading(doc, '1. 서론', 1)
    add_p(doc, '국내 금융회사는 전자금융거래의 안정성과 정보처리시스템 보호를 위해 업무용 시스템과 전산실 내 정보처리시스템을 외부통신망으로부터 분리·차단하는 망분리 체계를 운영해 왔다. 최근 금융당국은 SaaS와 생성형 AI 활용을 확대하기 위해 망분리 규제를 단계적으로 개선하면서도, 규제 완화가 보안수준 저하로 이어지지 않도록 금융회사의 위험평가와 자율적 보안통제를 강화하는 방향을 제시하고 있다[6][7]. 이에 따라 연구의 초점은 망분리의 단순 유지 또는 폐지가 아니라, 기존 망분리가 담당하던 외부위협 차단, 내부 확산 억제 및 정보반출 제한 기능을 어떤 세밀한 통제로 보완할 것인지로 이동하고 있다.')
    add_p(doc, '국가 망 보안체계(N2SF)는 업무정보와 정보시스템을 C(Classified), S(Sensitive), O(Open) 등급으로 구분하고, 정보서비스를 위치·주체·객체로 모델링하여 정보 생산·저장과 정보 이동 원칙을 적용한 후 필요한 보안통제를 선택하도록 한다[1]. 부록 2의 정보서비스 모델 해설서는 정보서비스 개요, 구성요소 분석, 위치·주체·객체 모델링 및 C/S/O 평가, 보안원칙 적용, 보안위협 식별, 보안 요구사항과 통제항목 선정의 순서로 참조모델을 제시한다[2]-[5].')
    add_p(doc, '다만 금융권 N2SF 적용 연구를 하나의 논문에서 모델 제안과 구현 검증까지 모두 다루면, 개념적 기여와 실증적 기여의 경계가 불명확해질 수 있다. 본 연구는 두 개의 논문 트랙을 명시적으로 구분한다. Track 1은 금융권의 망·시스템 위치 중심 운영환경을 고려하여 시스템·업무별 N2SF 모델을 제안하고, Track 2인 본 연구는 Track 1의 모델을 전제로 구체적인 보안통제를 도출하고 실험환경에서 검증한다. 따라서 본 논문은 시스템·업무별 모델 자체의 타당성을 다시 주장하지 않으며, 모델을 기술적 통제로 구현했을 때 통신경로와 정보이동 결과가 어떻게 달라지는지를 분석한다.')
    add_figure(doc, 'fig1_track_relationship.png', '그림 1. 두 논문 트랙의 역할 분리', 15.5)
    add_p(doc, '본 연구의 적용대상은 가상의 금융기관 A이다. 기관 A는 고객정보 관리, 여신심사, 신용평가, 자금세탁방지, 업무승인 및 공개·외부서비스 기능을 보유한다고 가정한다. 내부 업무와 개인정보·개인신용정보를 처리하는 업무영역은 S등급으로, 승인된 공개정보와 외부서비스 영역은 O등급으로 설정한다. C등급 정보는 연구범위에서 제외하며 실제 운영 중 C등급 정보가 확인되는 경우 별도의 상위 보호영역으로 분리해야 한다.')
    add_p(doc, '본 연구는 NSDI 2025의 ZTS가 제안한 역할 기반 마이크로세그멘테이션과 통신 그래프 기반 정책 평가 절차를 채택한다[8]. 단, ZTS의 역할 자동추론 알고리즘 전체를 재현하는 것이 아니라, Track 1에서 이미 정의된 시스템·업무 역할을 정답 레이블로 사용하고 역할별 허용 통신 그래프를 정책으로 코드화한다. 비교군과 제안군에 동일한 정상·비인가 흐름을 발생시켜 정책 위반률, 비인가 흐름 성공률, 직접 도달 가능한 데이터베이스 수, S/O 전송통제 효과 및 처리지연을 측정한다.')
    add_p(doc, '본 연구의 기여는 다음과 같다.', first_indent=False)
    add_bullet(doc, 'Track 1의 금융권 N2SF 시스템·업무별 모델을 전제로 한 후속 검증범위를 명확히 정의한다.')
    add_bullet(doc, 'N2SF 부록 2의 절차에 따라 가상 금융기관 A의 보안위협, 요구사항 및 통제항목을 도출한다.')
    add_bullet(doc, '업무별 컨테이너 네트워크, PEP/PDP 및 Transfer CDS를 이용한 재현 가능한 실험환경을 제시한다.')
    add_bullet(doc, '결과를 사전에 가정하지 않고 측정지표, 통계방법 및 빈 결과표를 제공하여 실제 실험 후 채울 수 있도록 한다.')

    # 2 background
    add_heading(doc, '2. 연구 배경 및 검증 방법론 선정', 1)
    add_heading(doc, '2.1 N2SF 정보서비스 모델 해설서의 통제 도출 구조', 2)
    add_p(doc, 'N2SF 부록 2는 공통 정보서비스를 대상으로 상위 수준의 참조구조와 보안대책 수립방법을 제시한다. 각 기관은 제시된 통제항목을 절대적 체크리스트로 적용하는 것이 아니라, 기관의 정보서비스 구성과 위협에 따라 조정하거나 추가해야 한다[2]-[5]. 본 연구는 이러한 성격을 반영하여 부록 2를 직접 복제하지 않고, 금융기관 A의 업무시스템과 데이터베이스, 업무 간 API 호출, S/O 경계 전송에 맞게 통제항목을 재구성한다.')
    add_figure(doc, 'fig5_appendix_process.png', '그림 2. N2SF 부록 2 형식에 따른 통제 도출 절차', 15.5)
    add_p(doc, '특히 부록 2-11은 CDS를 단순한 네트워크 연결 장비가 아니라 서로 다른 보안정책 영역 사이에서 승인, 검증, 무해화, 추적 및 선택적 전달을 수행하는 정책 기반 정보통제시스템으로 설명한다[5]. 이에 따라 본 연구의 S/O 경계는 단순 방화벽 포트 허용이 아니라 전송 주체, 목적, 정보등급, 승인상태 및 콘텐츠 검증결과를 함께 평가하는 Transfer CDS로 구현한다.')

    add_heading(doc, '2.2 검토한 국외 탑티어 방법론', 2)
    methods = [
        ('Header Space Analysis', 'NSDI 2012', '패킷 헤더와 네트워크 변환을 모델링하여 도달성, 루프, 격리 및 누출을 정적으로 검증[9]', '정책 배포 전 네트워크 불변조건 검증에 적합', 'Docker 기반 애플리케이션 문맥·정보등급 평가를 별도 구현해야 함'),
        ('NetPlumber', 'NSDI 2013', '규칙 의존성 그래프를 유지하여 네트워크 변경 시 정책 준수여부를 증분 검증[10]', '정책 변경 검증과 운영 자동화에 적합', 'SDN·포워딩 규칙 중심이므로 업무역할 및 데이터 목적 통제와 직접 대응이 약함'),
        ('Zanzibar', 'USENIX ATC 2019', '다양한 서비스의 객체 권한을 일관된 데이터 모델과 관계 기반 정책으로 평가[11]', '사용자-역할-객체 권한관계 검증에 적합', '네트워크 수평 이동과 S/O 경계의 실제 도달성 측정은 별도 필요'),
        ('NetVigil', 'NSDI 2024', '동서 트래픽 흐름 로그의 그래프 특징과 학습기법을 이용한 이상탐지[12]', '정상 패턴에서 벗어난 동적 행위 탐지에 적합', '학습데이터와 공격데이터 규모가 필요하고 본 연구의 소규모 환경에는 과도함'),
        ('ZTS', 'NSDI 2025', '흐름 텔레메트리에서 역할과 통신 그래프를 구성하고 역할 수준 마이크로세그멘테이션 정책을 적용[8]', '시스템·업무별 경계와 수평 이동 검증에 직접 대응', '원 논문의 역할 자동추론 전체를 재현하려면 대규모 운영데이터가 필요')
    ]
    add_table(doc, ['방법론','발표처','핵심 개념','적합성','본 연구 적용 한계'], methods, widths=[2.4,2.2,4.5,3.4,4.0], font_size=7.6)
    add_caption(doc, '표 1. 국외 탑티어 방법론 비교')

    add_heading(doc, '2.3 방법론 선택: ZTS의 역할 기반 마이크로세그멘테이션 적용', 2)
    add_p(doc, '본 연구는 ZTS를 우선 방법론으로 선택한다. 첫째, Track 1의 핵심 산출물이 시스템과 업무기능의 경계이므로, 역할 수준의 통신정책으로 변환하기 쉽다. 둘째, 서비스 간 통신 그래프와 정책 위반률은 Docker 기반 소규모 실험환경에서도 재현 가능하다. 셋째, 비인가 수평 이동, 업무 간 DB 직접접속, 경계 우회와 같은 금융권 망분리 완화의 핵심 위험을 정량화할 수 있다.')
    add_p(doc, '그러나 본 연구는 ZTS의 기계학습 기반 역할 추론을 그대로 적용하지 않는다. Track 1에서 업무 역할과 시스템 경계가 이미 제안되었다는 전제이므로, 고객·여신·신용·AML·승인 역할을 사전 정의하고 이를 정책 레이블로 사용한다. 통신 텔레메트리로부터 역할을 추론하는 대신, 승인된 업무흐름과 관측 흐름을 비교하여 정책 위반 여부를 평가한다. 이러한 연구설계는 ZTS의 전체 성능을 재현하는 연구가 아니라, 그 논문의 역할 기반 세분화와 평가절차를 금융권 N2SF에 적용한 실증연구로 한정된다.')
    add_figure(doc, 'fig4_method_pipeline.png', '그림 3. ZTS 절차의 금융권 N2SF 적용', 16.0)

    # 3 control derivation
    add_heading(doc, '3. 가상 금융기관 A의 보안위협 및 보안통제 도출', 1)
    add_heading(doc, '3.1 적용범위와 정보서비스 개요', 2)
    add_p(doc, '가상 금융기관 A는 은행 업무를 단순화한 연구용 조직으로 정의한다. 고객정보 업무는 고객기본정보와 계좌관계의 조회·갱신, 여신심사 업무는 신청접수와 심사흐름 관리, 신용평가 업무는 위험지표 산출, AML 업무는 이상거래 및 고객확인 검토, 승인 업무는 고위험 처리와 대외제공 승인 기능을 수행한다. 공개 API와 외부 생성형 AI 모의 서비스는 O영역에 배치한다.')
    components = [
        ('고객정보 업무','customer_app/customer_db','S','고객·계좌 관련 합성 레코드 조회·처리'),
        ('여신심사 업무','loan_app/loan_db','S','대출신청·심사상태·업무흐름 처리'),
        ('신용평가 업무','credit_app/credit_db','S','신용위험 평가 및 결과 제공'),
        ('자금세탁방지 업무','aml_app/aml_db','S','고객확인 및 이상거래 검토'),
        ('승인 업무','approval_app/approval_db','S','고위험 행위와 공개·외부제공 승인'),
        ('정책집행지점','PEP/API Gateway','S','업무 간 호출 중계 및 정책결정 요청'),
        ('정책결정지점','OPA(PDP)','S','사용자·역할·단말·목적·경로 평가'),
        ('S/O 연계체계','Transfer CDS','S/O 경계','등급·승인·콘텐츠·무결성 검증'),
        ('공개 서비스','public_app','O','승인된 O 파생본 제공'),
        ('외부 AI','external_ai','O','승인된 O 자료 활용 모의 서비스')
    ]
    add_table(doc, ['구분','구성요소','등급','주요 기능'], components, widths=[3.0,3.8,1.6,7.3], font_size=8)
    add_caption(doc, '표 2. 가상 금융기관 A의 정보서비스 구성요소')

    add_heading(doc, '3.2 위치·주체·객체 모델링 및 S/O 평가', 2)
    add_p(doc, '부록 2-2는 S등급 기관 전산망의 이용자 단말에서 O등급 생성형 AI 서비스를 사용하는 구조를 위치(S)-주체(S)-객체(O)로 모델링한다[2]. 본 연구는 이를 업무 간 동서 통신과 S/O 전송으로 확장한다. 위치는 업무별 S 세그먼트, S/O 연계체계 및 O 서비스 영역으로 구분하고, 주체는 사용자·업무서비스·관리자·정책집행 구성요소로 구분한다. 객체는 합성 업무레코드, API 요청·응답, 공개 파생본, 외부 AI 요청 및 감사로그로 정의한다.')
    lso = [
        ('위치','업무별 S 세그먼트','S','고객·여신·신용·AML·승인 업무'),
        ('위치','O 서비스 영역','O','공개 API·외부 AI'),
        ('위치','Transfer CDS','S/O 경계','정보 이동 검증 및 중계'),
        ('주체','업무 사용자·업무서비스','S','역할과 현재 업무목적에 따라 접근'),
        ('주체','외부 서비스','O','승인된 O 자료만 수신'),
        ('객체','내부 업무레코드·처리결과','S','원본 직접 외부전송 금지'),
        ('객체','공개 승인 파생본','O','승인·검증 후 외부전송 허용'),
        ('객체','정책결정·감사로그','S','변조 방지 및 추적성 확보')
    ]
    add_table(doc, ['유형','대상','등급','평가 원칙'], lso, widths=[1.7,4.5,1.8,7.8], font_size=8)
    add_caption(doc, '표 3. 위치·주체·객체 모델링 및 S/O 평가')

    add_heading(doc, '3.3 보안원칙 적용과 보안위협 식별', 2)
    add_p(doc, '정보 생산·저장 원칙에 따라 S 업무서비스가 생성한 원본과 처리결과는 해당 S 업무영역 또는 동일·상위 수준의 통제환경에 저장한다. O 영역으로 제공할 필요가 있는 경우 원본을 직접 이동하지 않고 최소화·비식별·검토·승인을 거친 O 파생본을 생성한다. 정보 이동 원칙에 따라 S 업무 간 통신은 업무 목적상 승인된 경로로 제한하고, S/O 이동은 Transfer CDS를 경유하도록 한다.')
    threats = [
        ('TH-F1','평면적 내부망에서 침해된 업무서비스가 다른 업무 DB에 직접 접근','S 업무영역','수평 이동·대량정보 접근'),
        ('TH-F2','사용자 인증 후 업무 역할과 무관한 서비스 호출이 허용','PEP/업무서비스','과권한·권한오남용'),
        ('TH-F3','업무 변경 후 기존 통신경로와 권한이 잔존','IAM/정책','권한 누적·미회수'),
        ('TH-F4','PEP를 우회한 서비스·DB 직접접속','네트워크 경계','통제 무력화·감사 누락'),
        ('TH-F5','S 원본이 공개 API나 외부 AI로 직접 전송','S/O 경계','개인정보·업무정보 유출'),
        ('TH-F6','O 파생본에 민감패턴 또는 내부 메타데이터가 포함','Transfer CDS','비식별 실패·메타데이터 유출'),
        ('TH-F7','정책엔진 장애 시 허용으로 처리','PEP/PDP','Fail-open에 의한 비인가 접근'),
        ('TH-F8','허용·거부 결정과 전송내용의 추적정보가 불완전','로그·감사','사고 원인분석 불가'),
        ('TH-F9','업무별 세그먼트 정책 변경 오류','정책관리','정상업무 차단 또는 비인가 경로 허용'),
        ('TH-F10','세분화로 인한 처리지연 증가','PEP/PDP/CDS','업무 가용성 저하')
    ]
    add_table(doc, ['위협 ID','보안위협','대상','영향'], threats, widths=[1.5,7.0,3.0,4.2], font_size=7.8)
    add_caption(doc, '표 4. 가상 금융기관 A의 보안위협')

    add_heading(doc, '3.4 보안 요구사항 및 보안통제 항목', 2)
    reqs = [
        ('SR-01','업무별 직접접속 범위 최소화','각 업무서비스는 자신의 DB만 직접 접근하고 타 업무 접근은 PEP를 경유','N2SF-LP-1, SG-2, SG-4, DU-M1','TH-F1, F4'),
        ('SR-02','역할·업무경로 기반 접근결정','사용자·역할·단말신뢰·출발업무·도착업무·목적을 매 요청 평가','N2SF-DA-4, SG-M4, IF-1, IF-9','TH-F2, F3'),
        ('SR-03','비인가 경로 자동 차단','허용 통신행렬에 없는 경로는 기본거부하고 위반을 기록','N2SF-SG-M5, IF-6, IF-8, IF-M2','TH-F2, F4'),
        ('SR-04','S/O 전송 정책 강제','S 원본 직접전송을 금지하고 승인된 O 파생본만 허용','N2SF-IF-14, CD-1, CD-11, CD-M3, DT-1','TH-F5'),
        ('SR-05','전송 콘텐츠와 무결성 검증','민감패턴·형식·메타데이터를 검사하고 콘텐츠 해시를 기록','N2SF-CD-4, CD-6, CD-7, CD-10, DT-6','TH-F6'),
        ('SR-06','우회접속 검증','서비스·DB 직접 도달성 및 비인가 경로를 정기 점검','N2SF-CD-9, SG-M5, IF-M4','TH-F4, F9'),
        ('SR-07','장애 시 안전한 차단','정책결정 실패·미인증·미승인 상태에서는 정보교환을 중단','N2SF-DT-2, DA-2','TH-F7'),
        ('SR-08','전 구간 추적성','정책 입력, 결정, 사유, 처리시간, 전송 해시를 감사로그로 기록','N2SF-IF-M2, CD-M2, DU-M2','TH-F8'),
        ('SR-09','정책 변경 검증','정책 변경 전 정상·비인가 시나리오를 자동 재실행','N2SF-SG-M5, IF-M4','TH-F9'),
        ('SR-10','성능 영향 관리','PEP/PDP/CDS의 p50·p95 지연과 정상업무 성공률을 측정','기관 추가 통제','TH-F10')
    ]
    add_table(doc, ['요구사항','목적','구현 요구','주요 N2SF 통제','대응 위협'], reqs, widths=[1.4,2.6,6.2,3.8,2.0], font_size=7.3)
    add_caption(doc, '표 5. 보안 요구사항과 N2SF 통제항목 매핑')
    add_p(doc, '표 5의 통제항목은 N2SF 부록 1과 부록 2-2, 2-8, 2-11의 관련 항목을 바탕으로 선정하였다[2][4][5]. 부록 2의 통제항목은 절대적 최소기준이 아니라 기관 환경에 따라 조정할 수 있는 참조사항이므로, 처리지연 측정과 정책 변경 회귀시험은 본 실험의 특성에 맞춘 추가 통제로 포함하였다.')

    # 4 implementation
    add_heading(doc, '4. 역할 기반 마이크로세그멘테이션 실험환경 구현', 1)
    add_heading(doc, '4.1 구현 원칙과 서버 표현', 2)
    add_p(doc, '실험은 단일 물리 호스트에서 Docker Compose를 이용하여 수행한다. 서로 다른 호스트 포트만 부여한 프로세스는 독립 서버 또는 독립 보안영역을 의미하지 않는다. 포트는 접속지점을 구분할 뿐 프로세스, 파일시스템, 네트워크 이름공간 및 저장영역을 분리하지 않기 때문이다. 따라서 본 연구는 업무 애플리케이션과 PostgreSQL을 각각 별도 컨테이너로 실행하고, 제안군에서는 업무별 Docker bridge network를 별도로 구성한다[14]. 논문에서는 이를 ‘컨테이너 기반 논리 서버 인스턴스’로 표현하며 물리 서버와 동일한 보증수준을 갖는다고 주장하지 않는다.')
    add_p(doc, '각 PostgreSQL 컨테이너는 내부적으로 동일한 5432 포트를 사용한다. 컨테이너와 네트워크 이름공간이 분리되어 있으므로 포트 충돌은 발생하지 않는다. 호스트 포트는 실험자가 접근해야 하는 PEP(18080), Transfer CDS(18090), OPA(18181)에만 매핑한다. Docker 공식 문서에 따르면 동일 네트워크에 연결된 서비스는 서비스명으로 서로 탐색·접근할 수 있고, 서로 네트워크를 공유하지 않는 서비스는 기본적으로 직접 통신하지 못한다[14].')

    add_heading(doc, '4.2 비교군 구성', 2)
    add_p(doc, '비교군은 현재의 모든 금융기관 환경을 동일하게 재현하는 것이 아니라, 시스템·업무별 세분화가 적용되지 않은 평면적 내부영역을 실험적으로 단순화한 구조이다. 고객, 여신, 신용, AML 및 승인 애플리케이션과 각 데이터베이스가 하나의 flat_s 네트워크를 공유한다. 정책은 인증된 사용자와 신뢰 단말에 대해 S영역 서비스 접근을 광범위하게 허용하고, S/O 전송은 승인 여부만 확인한다. 따라서 침해된 업무서비스가 다른 DB의 이름과 포트를 알고 있을 경우 직접 연결할 수 있는지가 비교군의 핵심 측정대상이 된다.')
    add_figure(doc, 'fig2_baseline_architecture.png', '그림 4. 비교군 실험환경', 16.0)

    add_heading(doc, '4.3 제안군 구성', 2)
    add_p(doc, '제안군은 고객, 여신, 신용, AML 및 승인 업무마다 독립 네트워크를 구성한다. 각 네트워크에는 해당 업무 애플리케이션, 해당 DB 및 PEP만 연결한다. PEP는 모든 업무 세그먼트에 연결되어 승인된 통신을 중계하지만, 개별 업무서비스는 다른 업무서비스나 DB를 직접 탐색할 수 없다. OPA는 별도의 통제 네트워크에 배치하며, PEP는 요청마다 사용자, 역할, 단말 신뢰상태, 출발업무, 도착업무, HTTP 행위, 목적 및 정보등급을 전달한다.')
    add_p(doc, '업무 간 허용 통신행렬은 Track 1의 업무흐름을 정책으로 변환한 것이다. 예를 들어 고객업무에서 여신업무로의 신청 전달, 여신업무에서 신용·AML·승인업무로의 검토 요청은 허용하지만, 고객업무에서 AML DB로의 직접접속과 승인업무에서 신용 DB로의 직접접속은 허용하지 않는다. PEP는 정책결정이 없거나 OPA 오류가 발생하면 요청을 차단하는 fail-closed 방식으로 동작한다.')
    add_figure(doc, 'fig3_proposed_architecture.png', '그림 5. 제안군 실험환경', 16.3)

    add_heading(doc, '4.4 정책집행과 Transfer CDS 구현', 2)
    add_p(doc, '정책결정은 Open Policy Agent와 Rego를 사용한다. OPA는 애플리케이션에서 정책결정을 분리하여 구조화된 요청문맥을 선언형 정책으로 평가할 수 있는 범용 정책엔진이다[13]. 본 환경의 PEP는 HTTP 요청을 받으면 OPA에 정책질의를 수행하고, 허용된 요청만 목적지 업무서비스로 전달한다. 허용·거부, 거부사유 및 정책결정시간은 JSONL 감사로그로 기록한다.')
    add_p(doc, 'Transfer CDS는 부록 2-11의 Transfer CDS 개념을 연구용으로 단순화한 구성요소이다[5]. 요청자의 사용자·역할·단말 신뢰상태, 목적지, 정보등급, 승인상태, 콘텐츠 안전성 및 원본 객체 ID를 평가한다. S 원본의 O영역 직접전송은 차단하고, 승인된 O 파생본만 공개 API 또는 외부 AI로 전달한다. 콘텐츠 검사는 실제 DLP·백신·CDR을 대체하지 않으며, 실험에서는 주민등록번호·카드번호 형태와 민감 키워드의 정규식 검출로 통제 흐름만 재현한다. 전송이 허용되면 SHA-256 해시와 결정근거를 기록한다.')
    impl = [
        ('업무서비스','FastAPI','업무별 API 및 자기 DB 접근'),
        ('데이터베이스','PostgreSQL 컨테이너 5개','업무별 합성 레코드 저장'),
        ('PEP','FastAPI API Gateway','업무 간 호출 중계·거부·감사'),
        ('PDP','Open Policy Agent/Rego','역할·경로·목적·등급 정책결정'),
        ('Transfer CDS','FastAPI + OPA','S/O 전송 승인·콘텐츠·무결성 검증'),
        ('네트워크','Docker Compose custom network','비교군 flat_s / 제안군 업무별 segment'),
        ('측정도구','Python/httpx/pandas/scipy','기능·도달성·지연·통계 분석')
    ]
    add_table(doc, ['구성요소','구현기술','실험상 역할'], impl, widths=[3.0,4.5,8.0], font_size=8)
    add_caption(doc, '표 6. 실험환경 구현기술')

    # 5 evaluation
    add_heading(doc, '5. 실험 설계 및 평가방법', 1)
    add_heading(doc, '5.1 연구질문과 가설', 2)
    rqs = [
        ('RQ1','업무별 역할정책은 정상 업무흐름을 유지하면서 비인가 업무 간 통신을 차단하는가?','H1: 제안군의 비인가 흐름 성공률은 비교군보다 낮다.'),
        ('RQ2','업무별 네트워크 분리는 침해된 서비스의 수평 이동과 DB 도달범위를 줄이는가?','H2: 제안군의 교차업무 DB 직접 도달률과 blast radius는 비교군보다 낮다.'),
        ('RQ3','Transfer CDS 통제는 S 원본의 O영역 전송을 차단하고 승인된 O 파생본은 허용하는가?','H3: 제안군은 모든 S 원본 전송을 차단하면서 승인된 O 파생본의 성공률을 유지한다.'),
        ('RQ4','세분화와 정책결정이 정상업무의 성능과 가용성에 미치는 영향은 어느 정도인가?','H4: 정상 흐름 성공률의 감소는 사전 정의 비열등성 한계 [δ 입력] 이내이고 p95 추가지연은 [예산 입력] ms 이내이다.')
    ]
    add_table(doc, ['연구질문','내용','검증가설'], rqs, widths=[1.4,7.2,7.2], font_size=7.8)
    add_caption(doc, '표 7. 연구질문 및 가설')

    add_heading(doc, '5.2 독립변수·종속변수 및 통제조건', 2)
    add_p(doc, '독립변수는 네트워크 및 정책구성 방식으로, 비교군과 제안군의 두 수준을 갖는다. 종속변수는 정상·비인가 흐름의 허용결과, DB 직접 도달성, S/O 전송결과, 정책결정 및 요청 지연, 로그 생성여부이다. 애플리케이션 코드, 데이터베이스 스키마, 합성 데이터, 실행 호스트, Docker·Python 버전, 요청 시나리오와 반복횟수는 두 환경에서 동일하게 유지한다.')
    variables = [
        ('독립변수','환경 유형','비교군(flat network·광범위 정책) / 제안군(업무별 segment·역할정책·CDS)'),
        ('종속변수','보안효과','비인가 흐름 성공률, DB 도달률, blast radius, S/O 유출 성공률'),
        ('종속변수','가용성·성능','정상 흐름 성공률, p50·p95 지연, 처리실패율'),
        ('종속변수','운영성','정책 위반률, 감사로그 완전성, 정책 규칙 수·변경량'),
        ('통제변수','환경','동일 호스트, 동일 컨테이너 자원한도, 동일 이미지·데이터·요청순서'),
        ('통제변수','반복','보안 시나리오 30회 이상, 성능 요청 1,000회 이상, 동일 워밍업')
    ]
    add_table(doc, ['구분','변수','정의'], variables, widths=[2.3,3.2,10.3], font_size=8)
    add_caption(doc, '표 8. 실험변수와 통제조건')

    add_heading(doc, '5.3 실험 시나리오', 2)
    add_p(doc, '정상 업무흐름은 고객→여신, 여신→신용, 여신→AML, 여신→승인, 신용→승인, 승인→여신으로 구성한다. 비인가 흐름은 고객→AML, 고객→승인, AML→고객, 승인→신용, 여신→고객, 비신뢰 단말 요청, 미등록 출발서비스 및 허용되지 않은 HTTP 행위로 구성한다. 네트워크 도달성 시험에서는 각 업무 애플리케이션 컨테이너에서 자신의 DB와 다른 업무 DB의 5432 포트에 TCP 연결을 시도한다.')
    scenarios = [
        ('정상업무','A01-A06','승인된 업무 간 API 호출','비교군·제안군 모두 허용 필요'),
        ('비인가 경로','U01-U08','업무 역할 외 호출, 비신뢰 단말, 미등록 출발지, 비허용 행위','제안군 차단 여부'),
        ('직접 DB 접근','R01-R12','업무서비스에서 자기·타 업무 DB 포트 연결','교차업무 도달성 비교'),
        ('S/O 전송','C01-C06','O 승인본, S 원본, 민감패턴, 미승인, 비신뢰 단말','CDS 정책효과'),
        ('성능','P01','여신→신용 승인 흐름 반복','p50·p95 및 실패율'),
        ('정책 장애','F01','OPA 중지 또는 오류 응답','fail-closed 확인')
    ]
    add_table(doc, ['유형','ID','내용','평가목적'], scenarios, widths=[2.2,1.8,7.4,4.4], font_size=8)
    add_caption(doc, '표 9. 실험 시나리오')

    add_heading(doc, '5.4 평가 지표', 2)
    add_p(doc, 'ZTS의 정책 위반률 개념을 적용하여 승인 통신 그래프에 존재하지 않는 관측 통신 간선의 비율을 계산한다. 본 연구에서는 반복 요청을 모두 간선으로 계산하는 이벤트 기반 지표와, 중복을 제거한 고유 간선 기반 지표를 함께 보고한다.')
    add_equation(doc, 'Policy Violation Rate (PVR) = |E_observed - E_allowed| / |E_observed|')
    add_p(doc, '비인가 흐름 성공률은 비인가 시도 중 실제로 목적지에 도달한 비율이며, 차단률은 그 보수이다. 정상 흐름 성공률은 승인된 시도 중 2xx 응답을 받은 비율로 정의한다.')
    add_equation(doc, 'Unauthorized Flow Success Rate = N_unauthorized,success / N_unauthorized,total')
    add_equation(doc, 'Block Rate = 1 - Unauthorized Flow Success Rate')
    add_p(doc, 'Blast radius는 하나의 업무서비스가 침해되었다고 가정했을 때 직접 도달할 수 있는 다른 업무서비스와 DB의 수로 정의한다. 네트워크 DNS 해석과 TCP 연결의 성공여부를 기준으로 계산하되, PEP를 통한 승인된 업무호출은 직접 도달성에서 제외한다.')
    add_equation(doc, 'BR(v) = |{u ∈ V : direct_reachable(v,u)=1, u≠v}|')
    add_p(doc, 'S/O 전송통제는 S 원본 전송 성공률, 승인된 O 파생본 성공률 및 민감패턴 포함 O 자료의 차단률로 평가한다. 성능은 워밍업 이후 성공 요청의 p50, p95, 평균 및 IQR을 제시하고, 제안군과 비교군의 차이를 계산한다. 감사로그 완전성은 모든 정책결정 시도 중 필수 필드를 포함한 로그가 생성된 비율로 측정한다.')
    metrics = [
        ('정상 흐름 성공률','정상 요청의 2xx 비율','기능 가용성'),
        ('정책 위반률(PVR)','허용 그래프에 없는 관측 간선 비율','정책 적합성'),
        ('비인가 흐름 성공/차단률','비인가 요청의 도달/차단 비율','접근통제 효과'),
        ('교차업무 DB 도달률','타 업무 DB TCP 연결 성공 비율','수평 이동 위험'),
        ('Blast radius','침해 노드에서 직접 도달 가능한 타 자산 수','공격 확산 범위'),
        ('S/O 유출 성공률','S 원본의 O영역 전달 성공 비율','정보이동 통제'),
        ('O 승인본 성공률','승인 O 파생본 전달 성공 비율','업무 활용성'),
        ('p50·p95 지연','정상 요청 응답시간 분포','성능 오버헤드'),
        ('감사로그 완전성','필수 필드를 갖춘 로그 비율','추적성')
    ]
    add_table(doc, ['지표','산정방법','의미'], metrics, widths=[4.0,7.0,4.8], font_size=8)
    add_caption(doc, '표 10. 평가 지표')

    add_heading(doc, '5.5 반복측정과 통계 분석', 2)
    add_p(doc, '기능·보안 시나리오는 각 환경에서 30회 이상 반복하고, 성능은 50회 이상 워밍업한 후 1,000회 이상 요청한다. 반복횟수, 유의수준, 비열등성 한계 및 허용 지연예산은 결과 확인 전에 고정한다. 비율 비교는 표본수가 작은 경우 Fisher의 정확검정을, 충분한 경우 카이제곱검정을 사용한다. 지연시간은 비정규성과 이상값을 고려하여 Mann-Whitney U 검정을 우선 적용하고 중앙값 차이의 95% 부트스트랩 신뢰구간을 함께 제시한다. 여러 시나리오군을 동시에 검정하는 경우 Holm 방식으로 다중비교를 보정한다.')
    add_p(doc, '실험 순서에 따른 캐시와 자원상태 영향을 줄이기 위해 비교군과 제안군의 실행순서를 번갈아 수행하거나, 각 환경을 독립적으로 재시작한 뒤 동일한 워밍업을 수행한다. 실험호스트의 CPU·메모리·운영체제·Docker 버전과 컨테이너 이미지 digest를 기록한다. 오류·누락 결과를 임의로 제거하지 않고 실패유형과 제외기준을 사전에 정의한다.')

    # Results
    add_heading(doc, '6. 실험 결과 작성 구조', 1)
    add_p(doc, '본 절은 실제 실험 후 측정값을 입력하기 위한 구조이다. 현재 초안에는 가상의 수치, 예상 향상률 또는 임의의 통계결과를 기재하지 않는다. 측정 전에는 아래 표의 셀을 비워 두고, 실험 산출물 CSV와 분석코드의 결과만 전사한다.')
    add_heading(doc, '6.1 정상·비인가 업무 흐름', 2)
    rows = [
        ('정상 흐름 성공률','[입력]','[입력]','[입력]','[입력]'),
        ('비인가 흐름 성공률','[입력]','[입력]','[입력]','[입력]'),
        ('비인가 흐름 차단률','[입력]','[입력]','[입력]','[입력]'),
        ('정책 위반률(이벤트)','[입력]','[입력]','[입력]','[입력]'),
        ('정책 위반률(고유 간선)','[입력]','[입력]','[입력]','[입력]')
    ]
    add_table(doc, ['지표','비교군','제안군','차이/효과크기','p-value/95% CI'], rows, widths=[4.3,2.7,2.7,3.2,3.0], font_size=8)
    add_caption(doc, '표 11. 정상 및 비인가 흐름 결과')
    add_placeholder(doc, 'security_effectiveness.png 삽입 및 결과 해석')

    add_heading(doc, '6.2 수평 이동과 공격 확산 범위', 2)
    rows = [
        ('자기 DB 도달률','[입력]','[입력]','정상 연결 유지 여부'),
        ('교차업무 DB 직접 도달률','[입력]','[입력]','감소율 [입력]'),
        ('평균 blast radius','[입력]','[입력]','차이 [입력]'),
        ('최대 blast radius','[입력]','[입력]','차이 [입력]')
    ]
    add_table(doc, ['지표','비교군','제안군','해석'], rows, widths=[5.0,3.0,3.0,4.8], font_size=8)
    add_caption(doc, '표 12. 네트워크 도달성 및 blast radius 결과')
    add_placeholder(doc, '업무별 도달성 행렬 또는 히트맵 삽입')

    add_heading(doc, '6.3 S/O 정보이동 통제', 2)
    rows = [
        ('승인된 O 파생본 전송 성공률','[입력]','[입력]','[입력]'),
        ('S 원본 외부전송 성공률','[입력]','[입력]','[입력]'),
        ('민감패턴 포함 자료 차단률','[입력]','[입력]','[입력]'),
        ('미승인 전송 차단률','[입력]','[입력]','[입력]'),
        ('CDS 감사로그 완전성','[입력]','[입력]','[입력]')
    ]
    add_table(doc, ['지표','비교군','제안군','해석'], rows, widths=[5.5,3.0,3.0,4.3], font_size=8)
    add_caption(doc, '표 13. S/O 정보이동 결과')

    add_heading(doc, '6.4 성능 및 운영 오버헤드', 2)
    rows = [
        ('정상 요청 성공건수','[입력]','[입력]','-'),
        ('p50 지연(ms)','[입력]','[입력]','차이 [입력]'),
        ('p95 지연(ms)','[입력]','[입력]','차이 [입력]'),
        ('IQR(ms)','[입력]','[입력]','-'),
        ('OPA 정책결정 p95(ms)','해당없음/입력','[입력]','-'),
        ('정책규칙 수','[입력]','[입력]','변경량 분석')
    ]
    add_table(doc, ['지표','비교군','제안군','효과/비고'], rows, widths=[5.0,3.1,3.1,4.6], font_size=8)
    add_caption(doc, '표 14. 성능 및 운영 오버헤드 결과')
    add_placeholder(doc, 'latency_boxplot.png 삽입 및 비열등성·지연예산 판단')

    # Discussion
    add_heading(doc, '7. 논의', 1)
    add_heading(doc, '7.1 결과 해석 원칙', 2)
    add_p(doc, '실험 후 결과는 단순한 향상률만으로 해석하지 않는다. 제안군에서 비인가 흐름과 교차업무 DB 도달성이 감소하더라도 정상업무 성공률이 유의하게 저하되거나 p95 지연이 사전 예산을 초과하면 운영 적용성이 제한된다. 반대로 성능 영향이 작더라도 정책 우회경로가 남거나 감사로그가 누락되면 N2SF 정보이동 통제를 충족했다고 판단할 수 없다. 따라서 보안효과, 업무가용성, 성능 및 추적성을 함께 평가한다.')
    add_p(doc, '정상 흐름 차단이 발견되면 Track 1의 업무흐름 정의가 불완전한지, Rego 정책이 과도하게 제한적인지, 또는 테스트 시나리오가 실제 승인업무를 잘못 표현했는지를 구분해야 한다. 비인가 흐름이 허용되면 네트워크 공유, PEP 우회, 정책 조건 누락, DNS·호스트 포트 노출, 관리용 인터페이스의 예외설정을 조사한다.')

    add_heading(doc, '7.2 N2SF 부록 2 형식의 학술적 활용', 2)
    add_p(doc, '본 연구는 N2SF 부록 2의 문서구조를 단순 요약하는 데 그치지 않고, 위협-요구사항-통제항목-구현-시험사례 간 추적성을 제공한다. 각 통제는 docs/control_mapping.csv에서 구현구성요소와 시험 ID에 연결된다. 이는 보안통제 항목의 존재 여부를 설문으로 확인하는 방식보다, 실제 정보흐름이 허용·차단되는지를 실행증적으로 확인할 수 있다는 장점이 있다.')

    add_heading(doc, '7.3 연구의 한계', 2)
    limits = [
        '컨테이너 네트워크는 실제 금융회사의 물리적 망분리, 가상화 플랫폼, 방화벽, NAC 및 망연계 솔루션의 보증수준을 재현하지 않는다.',
        '가상 금융기관 A의 업무흐름은 실험을 위한 단순화 모델이며 실제 금융회사별 조직·시스템·규정 차이를 대표하지 않는다.',
        '정규식 콘텐츠 검사는 실제 DLP, 백신, CDR, 문서포맷 검증 및 개인정보 탐지 성능을 평가하지 않는다.',
        'OPA 요청헤더는 실제 IdP, MFA, EDR, NAC, IAM/PAM에서 제공되는 신뢰신호를 추상화한다. 헤더 위조를 막기 위한 mTLS와 서명은 후속 구현이 필요하다.',
        '본 연구는 ZTS의 역할 자동추론 알고리즘 전체를 재현하지 않고 역할 기반 정책과 정책 위반률 평가를 적용한다.',
        '단일 호스트 실험의 지연결과는 실제 분산 금융망의 네트워크 지연과 처리량으로 일반화할 수 없다.'
    ]
    for x in limits: add_bullet(doc, x)

    add_heading(doc, '7.4 향후 연구', 2)
    add_p(doc, '후속 연구에서는 실제 기관의 익명화 NetFlow·API 로그 또는 현실적 합성 업무로그를 사용하여 허용 통신 그래프를 학습하고, 업무변경과 인사이동에 따른 정책 갱신비용을 분석할 필요가 있다. 또한 Header Space Analysis 또는 NetPlumber 계열의 정적·증분 검증을 결합하여 정책 배포 전 도달성 불변조건을 검사하고, NetVigil과 같은 동서 트래픽 이상탐지를 보완통제로 추가할 수 있다. 실제 금융권 적용에서는 OPA 입력 신호를 IAM, PAM, EDR, NAC, DLP, SIEM/SOAR와 연계하여 정책 결정과 사고대응의 자동화를 검증해야 한다.')

    # conclusion
    add_heading(doc, '8. 결론', 1)
    add_p(doc, '본 연구는 금융권 N2SF 시스템·업무별 모델을 제안하는 Track 1과 구분하여, 해당 모델을 보안통제로 구체화하고 실험적으로 검증하는 Track 2의 연구설계를 제시하였다. 가상 금융기관 A를 대상으로 N2SF 부록 2의 정보서비스 개요, 위치·주체·객체 모델링, 보안원칙, 보안위협, 요구사항 및 통제항목 도출절차를 적용하였다. 또한 ZTS의 역할 기반 마이크로세그멘테이션 개념을 채택하여 업무별 컨테이너 네트워크, PEP/PDP 및 Transfer CDS 기반 비교환경을 설계하였다.')
    add_p(doc, '본 초안에서는 어떠한 가상 결과도 제시하지 않았다. 정상·비인가 통신, DB 직접 도달성, S/O 전송, 정책 장애 및 성능 시나리오를 실제로 실행한 후, 생성된 CSV와 통계분석 결과를 제6장의 빈 표와 그림 위치에 입력해야 한다. 따라서 최종 논문의 결론은 실험 전 가설이 아니라 측정된 결과와 한계에 근거하여 작성되어야 한다.')

    # References
    add_heading(doc, '참고문헌', 1)
    refs = [
        '[1] 국가정보원·국가보안기술연구소, 「국가 망 보안체계 보안 가이드라인 1.0」, 2025.',
        '[2] 국가정보원·국가보안기술연구소, 「국가 망 보안체계 보안 가이드라인 정보서비스 모델 해설서: 모델 2. 업무환경에서 생성형 AI 활용」, 부록 2-2, 2025.',
        '[3] 국가정보원·국가보안기술연구소, 「국가 망 보안체계 보안 가이드라인 정보서비스 모델 해설서: 모델 3. 외부 클라우드 활용 업무협업 체계」, 부록 2-3, 2025.',
        '[4] 국가정보원·국가보안기술연구소, 「국가 망 보안체계 보안 가이드라인 정보서비스 모델 해설서: 모델 8. 클라우드 기반 통합문서체계」, 부록 2-8, 2025.',
        '[5] 국가정보원·국가보안기술연구소, 「국가 망 보안체계 보안 가이드라인 정보서비스 모델 해설서: 모델 11. 정보 연계를 위한 CDS 구성」, 부록 2-11, 2025.',
        '[6] 금융위원회, 「전자금융감독규정」 제15조, 시행 2026. 2. 13.',
        '[7] 금융위원회, 「금융분야 망분리 개선 로드맵」, 2024; 금융권 내부업무망 SaaS 규제 개선 관련 보도자료, 2026.',
        '[8] S. K. Mani et al., “Securing Public Cloud Networks with Efficient Role-based Micro-Segmentation,” 22nd USENIX Symposium on Networked Systems Design and Implementation (NSDI 2025), 2025.',
        '[9] P. Kazemian, G. Varghese, and N. McKeown, “Header Space Analysis: Static Checking for Networks,” 9th USENIX Symposium on Networked Systems Design and Implementation (NSDI 2012), pp. 113-126, 2012.',
        '[10] P. Kazemian et al., “Real Time Network Policy Checking Using Header Space Analysis,” 10th USENIX Symposium on Networked Systems Design and Implementation (NSDI 2013), pp. 99-111, 2013.',
        '[11] R. Pang et al., “Zanzibar: Google’s Consistent, Global Authorization System,” 2019 USENIX Annual Technical Conference, pp. 33-46, 2019.',
        '[12] K. Hsieh et al., “NetVigil: Robust and Low-Cost Anomaly Detection for East-West Data Center Security,” 21st USENIX Symposium on Networked Systems Design and Implementation (NSDI 2024), pp. 1771-1789, 2024.',
        '[13] Open Policy Agent, “OPA Documentation: Policy Language and REST API,” accessed 2026-08-04.',
        '[14] Docker, “Compose Networking and Compose File Networks Reference,” accessed 2026-08-04.',
        '[15] PostgreSQL Docker Community, “Postgres Docker Official Image,” accessed 2026-08-04.'
    ]
    for ref in refs:
        p=doc.add_paragraph(ref)
        p.paragraph_format.left_indent=Pt(12); p.paragraph_format.first_line_indent=Pt(-12)
        p.paragraph_format.space_after=Pt(3); p.paragraph_format.line_spacing=1.15

    # Appendix
    add_heading(doc, '부록 A. 실험 실행 절차', 1)
    add_heading(doc, 'A.1 패키지 구성', 2)
    add_p(doc, '재현 패키지의 루트 디렉터리는 financial_n2sf_track2이며, 비교군·제안군 Docker Compose 파일, FastAPI 업무서비스, PEP, Transfer CDS, OPA/Rego 정책, PostgreSQL 초기화 스크립트, 시나리오 CSV, 실행·분석 스크립트 및 결과 템플릿을 포함한다.')
    files = [
        ('compose.baseline.yml','평면적 S 네트워크 비교군'),
        ('compose.proposed.yml','업무별 마이크로세그먼트 제안군'),
        ('policies/*.rego','비교군·제안군 접근정책과 CDS 정책'),
        ('scenarios/*.csv','정상·비인가·CDS·도달성 시나리오'),
        ('scripts/run_*.py','정책·CDS·도달성·성능 시험'),
        ('scripts/analyze_results.py','요약·통계·그래프 생성'),
        ('docs/control_mapping.csv','N2SF 통제-구현-시험 추적표'),
        ('results/templates/*.csv','실험 전 빈 결과형식')
    ]
    add_table(doc, ['파일/경로','용도'], files, widths=[6.2,9.6], font_size=8)
    add_caption(doc, '표 A-1. 재현 패키지 구성')

    add_heading(doc, 'A.2 실행 순서', 2)
    commands = [
        'python -m venv .venv',
        'source .venv/bin/activate  # Windows: .venv\\Scripts\\activate',
        'pip install -r requirements.txt',
        'docker compose -f compose.baseline.yml up -d --build',
        './scripts/run_all.sh baseline',
        'docker compose -f compose.baseline.yml down -v',
        'docker compose -f compose.proposed.yml up -d --build',
        './scripts/run_all.sh proposed',
        'docker compose -f compose.proposed.yml down -v',
        'python scripts/analyze_results.py --results-dir results'
    ]
    for c in commands:
        p=doc.add_paragraph()
        p.paragraph_format.left_indent=Pt(12); p.paragraph_format.space_after=Pt(2)
        r=p.add_run(c); r.font.name='DejaVu Sans Mono'; r.font.size=Pt(8.4)

    add_heading(doc, 'A.3 결과 입력 체크리스트', 2)
    checklist = [
        '비교군과 제안군을 동일 호스트·동일 버전·동일 시나리오로 실행했는가?',
        '실행 전 반복횟수, 유의수준, 비열등성 한계 및 허용 지연예산을 고정했는가?',
        '보안 시나리오 실패와 네트워크 오류를 임의로 제외하지 않았는가?',
        'raw CSV, 정책파일, 이미지 digest, 호스트 사양을 보존했는가?',
        '제6장의 모든 수치가 분석 산출물과 일치하는가?',
        '향상되지 않은 결과와 정상업무 차단 사례도 함께 보고했는가?',
        'Docker 기반 논리적 격리를 물리적 망분리와 동일하다고 과장하지 않았는가?'
    ]
    for item in checklist: add_bullet(doc, '□ ' + item)

    add_heading(doc, '부록 B. 결과 작성용 서술 템플릿', 1)
    add_p(doc, '다음 문장은 실제 결과가 생성된 후 수치와 통계결과를 입력하여 사용한다. 결과가 가설과 다를 경우 문장 구조를 억지로 유지하지 말고 관측결과에 맞게 수정한다.')
    templates = [
        '비인가 업무 흐름 성공률은 비교군 [ ]에서 제안군 [ ]로 [증가/감소]하였으며, Fisher의 정확검정 결과 [통계량 및 p-value]로 나타났다.',
        '교차업무 DB 직접 도달률은 비교군 [ ]와 제안군 [ ]로 측정되었고, 업무별 네트워크 분리가 [어떠한 범위에서] 직접 접근경로를 제한하였다.',
        '승인된 O 파생본의 전송 성공률은 [ ]이었으며, S 원본 및 민감패턴 포함 자료의 차단률은 각각 [ ], [ ]로 측정되었다.',
        '제안군의 p95 지연은 비교군보다 [ ]ms [증가/감소]하였으며, 사전 정의한 허용예산 [ ]ms를 [충족/초과]하였다.',
        '따라서 H1-H4 중 [ ]는 지지되었고 [ ]는 지지되지 않았다. 주요 원인은 [정책경로/네트워크 구성/성능 병목]으로 분석되었다.'
    ]
    for t in templates: add_bullet(doc, t)

    # metadata
    doc.core_properties.title = TITLE
    doc.core_properties.subject = '금융권 N2SF Track 2 보안통제 및 실험설계'
    doc.core_properties.author = '[저자 입력]'
    doc.core_properties.keywords = 'N2SF, 금융권, 망분리, 마이크로세그멘테이션, OPA, CDS'
    doc.save(OUT)
    print(OUT)

if __name__ == '__main__':
    main()
